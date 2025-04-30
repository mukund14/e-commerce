# streamlit_app.py

import streamlit as st
import nbformat
import base64
import tempfile
from pptx import Presentation
from pptx.util import Inches
from openai import OpenAI
import os

# Setup OpenAI client
client = OpenAI(api_key=st.secrets.get("OPENAI_API_KEY") or os.getenv("OPENAI_API_KEY"))

def generate_slide_content(markdown_text):
    """Uses OpenAI GPT-4 to generate slide title, speaker notes, and visual suggestion."""
    prompt = f"""
    You are an expert AI presentation assistant. Given the following notebook content:

    {markdown_text}

    Generate:
    1. A concise slide title
    2. 2-3 sentence speaker notes explaining the content clearly
    3. A suggested type of visual (e.g., pie chart, scatter plot, bar graph, process diagram)

    Respond clearly and label each part (Title:, Speaker Notes:, Visual:).
    """

    response = client.chat.completions.create(
        model="gpt-4-1106-preview",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.4,
        max_tokens=400
    )

    output = response.choices[0].message.content.strip()
    sections = {"title": "", "speaker_notes": "", "visual": ""}
    
    for line in output.splitlines():
        if line.startswith("Title:"):
            sections["title"] = line.replace("Title:", "").strip()
        elif line.startswith("Speaker Notes:"):
            sections["speaker_notes"] = line.replace("Speaker Notes:", "").strip()
        elif line.startswith("Visual:"):
            sections["visual"] = line.replace("Visual:", "").strip()

    return sections

def create_pptx(slides_data):
    """Create a PowerPoint file from parsed notebook content."""
    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0]

    # Title slide
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Data Science Project Overview"
    slide.placeholders[1].text = "Generated from Jupyter Notebook"

    # Content slides
    for slide_data in slides_data:
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_data["title"]

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.text = f"Visual Suggestion: {slide_data['visual']}\n\n"
        p = tf.add_paragraph()
        p.text = slide_data["markdown"]


        notes_slide = slide.notes_slide
        notes_text_frame = notes_slide.notes_text_frame
        notes_text_frame.text = slide_data["speaker_notes"]

    return prs

# Streamlit App
st.title("Notebook to Presentation (with GPT Speaker Notes)")

uploaded_file = st.file_uploader("Upload a .ipynb file", type=["ipynb"])

if uploaded_file:
    notebook = nbformat.read(uploaded_file, as_version=4)
    slides_data = []

    for cell in notebook.cells:
        if cell.cell_type == "markdown":
            ai_content = generate_slide_content(cell.source)
            slides_data.append({
                "markdown": cell.source,
                "title": ai_content["title"],
                "speaker_notes": ai_content["speaker_notes"],
                "visual": ai_content["visual"]
            })

    if slides_data:
        st.success(f"Successfully processed {len(slides_data)} slides.")
        if st.button("Generate and Download PPTX"):
            prs = create_pptx(slides_data)
            tmp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
            prs.save(tmp_file.name)

            with open(tmp_file.name, "rb") as f:
                st.download_button(
                    label="Download Presentation",
                    data=f,
                    file_name="notebook_presentation.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
                )
